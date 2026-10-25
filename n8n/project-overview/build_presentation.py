"""Build the project walkthrough and add a dedicated cover to the research deck.
Requires python-pptx and Pillow. Screenshots are captured from the actual local n8n UI.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'n8n/output'
SHOTS = OUT / 'project-screenshots'
NAVY='10263E'; TEAL='007F83'; INK='18334B'; MUTED='526578'; LIGHT='EDF4F6'; WHITE='FFFFFF'; GOLD='E9B85A'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)

def rect(slide,x,y,w,h,color):
 s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(color);s.line.fill.background();return s

def text(slide,x,y,w,h,value,size=18,color=INK,bold=False):
 box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=box.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
  p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
  p.space_after=Pt(9)
 return box

def slide(title,sub=''):
 s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,13.333,7.5,WHITE)
 rect(s,.45,.42,.1,.3,TEAL);text(s,.72,.4,11,.3,'MAVENS  /  WEEK 3 PROJECT',11,TEAL,True)
 text(s,.65,.98,12,.65,title,30,INK,True)
 if sub:text(s,.67,1.72,12,.55,sub,15,MUTED)
 rect(s,.65,7.05,12,.012,'D7E2E7');text(s,.67,7.14,11,.2,'Amit Gandre  •  October 25, 2026  •  Public-source research automation',9,MUTED)
 text(s,12.05,7.12,.6,.25,str(len(prs.slides)),10,MUTED)
 return s

def note(s,value):s.notes_slide.notes_text_frame.text=value

def card(s,x,y,w,h,label,body):
 rect(s,x,y,w,h,LIGHT);rect(s,x,y,.045,h,TEAL)
 text(s,x+.18,y+.16,w-.36,.48,label,19,TEAL,True)
 text(s,x+.18,y+.78,w-.36,h-.9,body,17,INK)

def picture(s,name,x,y,w,h):
 path=SHOTS/name;im=Image.open(path);ratio=min(w/im.width,h/im.height)
 ww,hh=im.width*ratio,im.height*ratio
 s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

def cover(p,title,subtitle,tag):
 s=p.slides.add_slide(p.slide_layouts[6]);rect(s,0,0,13.333,7.5,NAVY)
 rect(s,9.7,0,3.633,7.5,TEAL);rect(s,.72,1.02,.8,.06,GOLD)
 text(s,.75,.58,10,.3,'MAVENS  /  WEEK 3 PROJECT',13,WHITE,True)
 text(s,.75,1.65,8.4,2.4,title,40,WHITE,True)
 text(s,.8,4.35,8.0,1.1,subtitle,21,'C5E6E9')
 text(s,.8,6.3,8,.6,'Amit Gandre\nOctober 25, 2026',16,WHITE)
 text(s,10.05,1.45,2.8,3.0,tag,25,WHITE,True)
 text(s,10.05,6.65,2.7,.4,'GOALS → DELIVERY',11,WHITE,True)
 return s

s=cover(prs,'Public-source IT &\nCybersecurity Research','Project goals, agents, architecture and a practical walkthrough','4 AGENTS\n1 ORCHESTRATOR\nREUSABLE MEMORY')
note(s,'Project overview of the implemented local n8n system. This deck is distinct from the Target research presentation. Screenshots show actual local workflow canvases captured October 25, 2026.')

s=slide('Project goals','Turn a company name into a repeatable, source-linked research briefing.')
for x,y,label,body in [(0.67,2.55,'01  Identify comparable companies','Find up to five revenue peers using supported financial evidence and sector checks.'),(6.82,2.55,'02  Surface recent signals','Review three calendar months of IT, cybersecurity, technology and leadership news.'),(0.67,4.55,'03  Add professional context','Research CEO, CISO, CTO and CIO candidates, public career history and LinkedIn links.'),(6.82,4.55,'04  Deliver and remember','Create an editable PowerPoint and retain dated research for comparison and recall.')]:card(s,x,y,5.83,1.8,label,body)
note(s,'Five peers is a target, not a guarantee. Unsupported target revenue blocks research. An eligible-peer shortfall stays explicit. Search excerpts remain preliminary evidence.')

s=slide('Architecture: from company input to retained briefing','Explicit workflow routing connects search, local inference, validation and reporting.')
boxes=[(.7,2.75,'COMPANY FORM','Name • website • sector'),(3.9,2.75,'AGENT 1','Revenue peers + checks'),(7.1,2.75,'AGENTS 2 → 3','Search → extract signals'),(10.3,2.75,'AGENT 4','Leadership + background')]
for x,y,l,b in boxes:
 rect(s,x,y,2.45,1.1,LIGHT);text(s,x+.12,y+.15,2.22,.3,l,16,TEAL,True);text(s,x+.12,y+.58,2.22,.4,b,13)
for x in [3.22,6.42,9.62]:text(s,x,3.04,.5,.4,'→',25,TEAL,True)
rect(s,2.0,4.65,9.3,1.15,NAVY);text(s,2.25,4.82,8.8,.35,'ORCHESTRATOR  →  EDITABLE POWERPOINT  +  RESEARCH JSON',19,WHITE,True)
text(s,2.25,5.3,8.8,.3,'Append-only Data Table memory  •  Prior-run comparison  •  Download and recall',14,WHITE)
text(s,.8,6.25,11.9,.45,'Runtime: local n8n + Ollama/Gemma  |  External evidence: native Google/SerpApi nodes',17,MUTED)
note(s,'Actual routing: Agent 1 inside Research Orchestrator; Agent 2 calls Agent 3 internally. Orchestrator recalls leadership memory and calls Agent 4 for missing/stale profiles; the current cache is 168 hours (seven days). It recalls previous research, compiles PPTX and appends a snapshot. No HTTP Request nodes in current production workflows.')

s=slide('All agents and their responsibilities','The orchestrator coordinates the four research stages and owns final output.')
rows=[('Agent 1','Revenue peers','Company + financial evidence','Up to five eligible peers; evidence gaps'),('Agent 2','Public web and news','Target + selected peers + date window','Dated headlines, links and evidence package'),('Agent 3','IT / cyber signal analysis','Supplied public evidence','Validated, attributed signals and citations'),('Agent 4','Leadership background','Company + saved profile context','Role candidates, career/education claims, links'),('Orchestrator','Coordinate, compile, retain','Peer decisions + signals + profiles','PPTX + research JSON + persistent snapshot')]
y=2.55
for i,(a,b,c,d) in enumerate(rows):
 rect(s,.68,y,11.98,.69,LIGHT if i%2==0 else 'F8FAFB');text(s,.82,y+.12,1.55,.45,a,16,TEAL,True);text(s,2.4,y+.12,2.5,.45,b,16,INK,True);text(s,4.95,y+.12,3.05,.45,c,14);text(s,8.1,y+.12,4.35,.5,d,14);y+=.79
note(s,'Agents are bounded workflows with defined inputs/outputs and local model extraction, not unrestricted autonomous browsers. Agent 1 resides in the main orchestrator. Test workflows are separate from production.')

s=slide('Agent 1: build the revenue peer set','Actual Research Orchestrator canvas; Agent 1 runs in the opening financial stages.')
picture(s,'orchestrator-canvas.png',.68,2.4,12,2.35)
for x,l,b in [(0.7,'INPUT','Company identity, annual report or retained revenue evidence.'),(4.85,'CHECKS','Quoted amount, year, currency, company attribution and range.'),(9.0,'OUTPUT','Ranked eligible peers and explicit evidence gaps.')]:card(s,x,5.05,3.64,1.65,l,b)
note(s,'Screenshot: actual local Research Orchestrator. Financial flow includes report upload/retained-report fallback, peer discovery, revenue extraction, normalization and FX validation. Current user-authorized automatic selection sets sources_verified=false. The node name Finalize Reviewed Peers is retained, but there is no interactive pre-approval form in the current production graph.')

for title,sub,name,cards,notes in [
('Agent 2: collect public web and news','Native Google News + Google Web return evidence with source URLs.','agent2-canvas.png',[('SEARCH','Current low-credit mode uses two core searches per company.'),('FILTER','Apply research dates, company matching, URL cleanup and deduplication.'),('HANDOFF','Package source evidence and send it to Agent 3.')],'Actual Agent 2 canvas. Current low-credit settings disable the extra local AI query planner and automatic search retries. Six companies require up to 12 base Agent 2 search requests; financial discovery and leadership searches are additional. Unknown dates and weak matches remain review candidates.'),
('Agent 3: extract supported IT / cyber signals','Local Gemma extracts structured claims; deterministic checks narrow what is retained.','agent3-canvas.png',[('EXTRACT','Leadership, contracts, cyber news and technology signals.'),('VALIDATE','Check company attribution, source quotes and date window.'),('LABEL','Separate hiring, planned activity and announced deployment.')],'Actual Cybersecurity Signals canvas, called by Agent 2. Local Ollama gemma4:12b with structured output. Guards reject unsupported company claims, generic advice and contract-value inferences. Job advertisements are hiring evidence, not proof of installed technology.'),
('Agent 4: research public leadership background','Discover CEO / CISO / CTO / CIO candidates and preserve uncertainty.','agent4-canvas.png',[('DISCOVER','Two combined leadership searches, then validate candidates.'),('ENRICH','One public LinkedIn search per distinct leader; up to four people.'),('RETAIN','Append profiles with dates, source links and missing/conflicting roles.')],'Actual low-credit Agent 4 canvas. At most six initial searches per company: two discovery plus up to four leader background searches; no automatic retries. Extract Deloitte/KPMG/EY employment and SEC-school education only when supported. The reduced search plan may leave these fields missing. Public LinkedIn snippets are not full-profile retrieval. Orchestrator reuses complete/partial profiles for up to seven days.'),
('Research memory: retain, compare and recall','Actual recall workflow; production research is retained in the native n8n Data Table.','memory-canvas.png',[('APPEND','Keep research snapshots, leadership profiles and PowerPoint bytes.'),('COMPARE','Mark prior evidence and newly observed sources across runs.'),('RECALL','Search by company; use a run ID to retrieve an older saved version.')],'Actual Research Memory - Recall canvas: form → company key → read saved records → downloadable JSON/PPTX. Preview and test record types remain separate. The UI returns 20 matches at a time. Workflow JSON exports do not include the memory database; preserve the n8n data volume and local memory export.')]:
 s=slide(title,sub);picture(s,name,.68,2.4,12,2.35)
 for x,(l,b) in zip([.7,4.85,9.0],cards):card(s,x,5.05,3.64,1.65,l,b)
 note(s,notes)

s=slide('How to run: use the existing local installation','Start with Research Orchestrator at localhost:5678.')
steps=[('1','Start services','Keep Docker/n8n and local Ollama running; confirm gemma4:12b is available.'),('2','Open the workflow','Open Research Orchestrator, refresh the editor, click Execute workflow and open its form.'),('3','Submit the company','Enter real company name, official website, sector and geography. Attach an annual report when available.'),('4','Let the stages complete','Peer checks → public search → signal extraction → leadership recall/research → presentation and memory.'),('5','Download your briefing','Open Download Briefing → Binary → presentation. Use Research Memory - Recall for saved versions.')]
for i,(num,l,b) in enumerate(steps):
 y=2.43+i*.81;rect(s,.75,y,.42,.42,TEAL);text(s,.86,y+.045,.3,.3,num,16,WHITE,True);text(s,1.45,y,3.05,.52,l,18,INK,True);text(s,4.55,y,7.8,.65,b,16)
note(s,'Current automatic peer decision is user-authorized and continues with explicit gaps; it does not mean original financial sources have been independently verified. Leadership retrieval uses seven-day cached matching profiles. Unsupported target revenue stops the pipeline and produces a review diagnostic. Executing native searches uses SerpApi allowance.')

s=slide('Example input: reproduce the Target walkthrough','The main form supplies the company identity and financial-comparison context.')
card(s,.7,2.5,5.83,3.65,'FORM VALUES','Company Name: Target\nOfficial Website: https://corporate.target.com\nIndustry Sector: Retail\nGeography: United States\nAnnual Report: upload if available')
card(s,6.82,2.5,5.83,3.65,'OPTIONAL CONTROLS','Annual Report URL\nMinimum / Maximum Revenue Ratio\nFX Rates JSON when supported conversion is needed\n\nNews window: three calendar months')
text(s,.85,6.43,11.5,.35,'Expected output: a dated briefing, explicit coverage gaps, source links and a retained research snapshot.',16,MUTED)
note(s,'These labels match the current orchestrator form. Uploaded or retained financial evidence still requires original-source verification. A new run may return different eligible peers, source counts, leadership candidates or findings; the archived Target production run is not a promised result for future runs.')

s=slide('How to set up from GitHub','Import portable workflows, configure local dependencies and reconnect credentials.')
card(s,.7,2.5,3.82,3.75,'01  GET THE PROJECT','Clone amitgandre/Mavens-Week3-AmitG.\n\nUse n8n/saved-workflows/all-workflows.json or import individual production workflows.')
card(s,4.75,2.5,3.82,3.75,'02  CONNECT SERVICES','Run n8n with persistent storage.\n\nInstall Ollama + gemma4:12b. Configure Ollama and SerpApi credentials; native search nodes must be available.')
card(s,8.8,2.5,3.82,3.75,'03  INITIALIZE & TEST','Inspect Research Memory - Setup and create/map the native Data Table.\n\nPublish internal-call workflows, remap references if needed, then run the local model check.')
text(s,.85,6.48,11.5,.3,'Portable imports are inactive and contain no credential references. Memory and generated files remain local.',15,MUTED)
note(s,'Repository: https://github.com/amitgandre/Mavens-Week3-AmitG. For n8n in Docker on Mac, the Ollama credential endpoint is http://host.docker.internal:11434. Internal call workflows are Agent2PublicResearch, Cybersecurity Signals and Agent4LeadershipBackground. Native Data Table IDs and workflow references may require remapping on another instance. Test/cached workflows are optional. Root .env.example and research.py are historical prototypes, not current setup entry points. Preserve the existing database/volume to migrate research history.')

s=slide('Deliverables: a usable research system','Outputs support both presentation and further evidence review.')
for x,y,l,b in [(0.7,2.55,'EDITABLE POWERPOINT','Target-final-briefing.pptx\nPeer comparison, selected signals, leadership and linked sources.'),(6.82,2.55,'COMPLETE RESEARCH','Target-final-briefing.json\nFull evidence, quotations, source registry, dates and gaps.'),(0.7,4.55,'REUSABLE WORKFLOWS','22 portable production/test workflow exports, agent source, checks and documentation.'),(6.82,4.55,'PERSISTENT HISTORY','Native memory retains dated research and deck bytes; recall retrieves saved versions.')]:card(s,x,y,5.83,1.8,l,b)
note(s,'PDF is a shareable export of the presentation. Full local backup contains 28 workflows, including six temporary Target production/recovery/compilation workflows. Portable set contains 22 reusable workflows. Raw backups, memory exports and generated output artifacts are excluded from GitHub. This project overview is the second requested deck.')

s=slide('Example delivery: the Target research briefing','Actual generated presentation pages, retained separately from this project walkthrough.')
picture(s,'target-slide-2.png',.65,2.45,6.0,3.4);picture(s,'target-slide-3.png',6.75,2.45,5.95,3.4)
for x,n,l in [(1.0,'6','Companies'),(4.1,'14','Extracted signals'),(7.2,'288','Source records'),(10.3,'13','Leaders researched')]:
 text(s,x,6.0,2,.5,n,29,TEAL,True);text(s,x,6.53,2.5,.3,l,14,MUTED)
note(s,'Images are actual exported Target deck pages: peer revenue/signal comparison and company snapshots. Target plus Lowe’s, Albertsons, Kroger, Home Depot and Publix Super Markets. Metrics are from retained production research-178. The dedicated cover requested by the user increases the research deck from eleven to twelve slides; counts of companies/signals/sources are unchanged. Source records include public search evidence and background records, not 288 fully reviewed articles.')

s=slide('Validation and operating limits','The system preserves evidence boundaries while making the work repeatable.')
card(s,.7,2.55,5.83,3.65,'CHECKED IN THIS PROJECT','Company attribution and quoted evidence\nStructured extraction and date handling\nPeer eligibility and unsupported-revenue stops\nAppend-only memory and prior-run recall\nPowerPoint opens, exports and fits slide bounds')
card(s,6.82,2.55,5.83,3.65,'WHAT THE READER MUST KNOW','Search snippets need original-source review.\nRole candidates may be stale or disputed.\nHiring is not evidence of deployment.\nSignal counts are not security-risk scores.\nLow-credit search limits coverage; missing means insufficient evidence.')
note(s,'Recorded tests live under each agent directory and n8n/orchestrator. Deterministic checks and selected live/cached regressions were completed; this is not a claim of exhaustive correctness or independently verified current corporate facts. Financial, search and leadership stages all use quota beyond the Agent 2 base query count. Public-source retrieval does not require LinkedIn login, and no full-profile scraping occurs.')

s=slide('Handoff: code, workflow entry points and files','The project is ready to inspect, run and extend from the saved implementation.')
card(s,.7,2.5,7.35,3.55,'GITHUB REPOSITORY','github.com/amitgandre/Mavens-Week3-AmitG\n\nStart: README.md\nImport: n8n/saved-workflows/README.md\nRun / recall: Research Orchestrator; Research Memory - Recall\nPresentation source: n8n/project-overview/build_presentation.py')
card(s,8.3,2.5,4.32,3.55,'TWO PRESENTATIONS','1  Target research briefing\nBusiness research results\n\n2  Project overview\nGoals, agents, architecture, screenshots, operation and deliverables')
link=text(s,.85,6.4,11.7,.4,'Open the repository →',18,TEAL,True)
link.text_frame.paragraphs[0].runs[0].hyperlink.address='https://github.com/amitgandre/Mavens-Week3-AmitG'
note(s,'Generated PowerPoint/PDF deliverables remain local and are linked in the chat. Rebuild this deck with python3 n8n/project-overview/build_presentation.py after capturing screenshots in n8n/output/project-screenshots. Screenshots are real n8n workflow canvases; editable architecture and process summaries are explanatory diagrams, not screenshots.')

OUT.mkdir(exist_ok=True)
prs.save(OUT/'Project-overview.pptx')

# Give the existing research briefing a dedicated title slide, without changing research content.
target_path=OUT/'Target-final-briefing.pptx'
target=Presentation(target_path)
first_text=' '.join(shape.text for shape in target.slides[0].shapes if shape.has_text_frame)
if 'TARGET RESEARCH BRIEFING' not in first_text:
 s=cover(target,'Target\nIT & Cybersecurity Briefing','Target and five revenue peers\nPublic-source research • July 25 – October 25, 2026','TARGET RESEARCH BRIEFING\n\n6 COMPANIES\n14 SIGNALS')
 s.notes_slide.notes_text_frame.text='Dedicated research cover. Prepared for Amit Gandre. Research date October 25, 2026. Search snippets and leadership candidates need original-source verification; revenue peers are provisional. See subsequent slides for findings and linked evidence.'
 ids=target.slides._sldIdLst;ids.insert(0,ids[-1]);target.save(target_path)
print(f'Project overview: {len(prs.slides)} slides; Target briefing: {len(target.slides)} slides. Both have dedicated covers.')
